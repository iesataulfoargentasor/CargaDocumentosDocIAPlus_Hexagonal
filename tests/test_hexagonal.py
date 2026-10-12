"""Pruebas de fronteras, contratos y fallos que no cubre el recorrido nominal."""
import ast
from copy import deepcopy
import io
import json
from pathlib import Path
import subprocess
import sys
from uuid import uuid4
import zipfile
import pytest
from botocore.exceptions import ClientError
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject
from app.application.dto import UploadCommand
from app.application.document_service import DocumentService
from app.adapters.persistence.memory_repository import MemoryRepository
from app.adapters.persistence.record_mapper import to_record, from_record
from app.adapters.persistence.s3_repository import S3Repository
from app.adapters.validation.file_validator import DocumentFileValidator
from app.domain.errors import PermissionDenied, StorageError, ValidationError
from app.domain.user import User
from conftest import FakeS3, login, upload

ROOT = Path(__file__).resolve().parents[1]
ACTOR = User('docente',frozenset({'read','upload'}))
META = dict(titulo='Programación de prueba',ensenanza='IABD',modulo='SBD',curso_academico='2026-2027',estado_declarado='final',procedencia='Departamento')

def command(content=b'Contenido de prueba',**metadata):
    return UploadCommand({**META,**metadata},io.BytesIO(content),'prueba.txt')

def original_key(repo,record):
    return to_record(record,repo.prefix)['archivo']['clave_s3']

@pytest.mark.parametrize('backend',['memory','s3'])
def test_port_contract(backend):
    repo = MemoryRepository() if backend=='memory' else S3Repository('test',client=FakeS3())
    catalog=json.loads((ROOT/'config/catalogo.json').read_text())
    service=DocumentService(repo,DocumentFileValidator(),catalog,'Centro',1000)
    record=service.create(ACTOR,command())
    assert service.get(ACTOR,record['id'])==record
    assert service.list(ACTOR,{'q':'PRUEBA'})==[record]
    assert service.list(ACTOR,{'modulo':'PIA'})==[]
    assert service.download(ACTOR,record['id']).content==b'Contenido de prueba'
    assert 'clave_s3' not in record['archivo']
    with pytest.raises(PermissionDenied):
        service.create(User('lector',frozenset({'read'})),command())
    with pytest.raises(PermissionDenied):
        service.list(User('sin-permiso',frozenset()))


def test_core_can_import_without_frameworks():
    code='import sys; import app.domain.document, app.application.document_service; assert not any(k in sys.modules for k in ("flask","boto3","pypdf","docx","werkzeug"))'
    subprocess.run([sys.executable,'-c',code],cwd=ROOT,check=True)
    for layer in ('domain','application'):
        for path in (ROOT/'app'/layer).glob('*.py'):
            tree=ast.parse(path.read_text())
            for node in ast.walk(tree):
                names=[a.name for a in node.names] if isinstance(node,ast.Import) else [node.module or ''] if isinstance(node,ast.ImportFrom) else []
                assert not any(name.startswith(('flask','boto','werkzeug','pypdf','docx','app.adapters','app.config')) for name in names),path


def test_test_config_ignores_env_and_never_creates_aws(monkeypatch,tmp_path):
    import boto3
    import app.config
    from app import create_app
    monkeypatch.setenv('SECRET_KEY','wrong')
    monkeypatch.setenv('AWS_PROFILE','missing-profile')
    def forbidden(*a,**kw): raise AssertionError('No debe leer .env ni crear sesión AWS')
    monkeypatch.setattr(app.config,'dotenv_values',forbidden)
    monkeypatch.setattr(boto3,'Session',forbidden)
    app=create_app({'TESTING':True,'SECRET_KEY':'s'*64,'USERS_FILE':str(tmp_path/'absent')},repository=MemoryRepository())
    assert app.test_client().get('/login').status_code==200
    assert app.config['AWS_OPTIONS']=={}


def test_raw_json_compatibility_and_stream_closure(env):
    _,c,s3,repo=env;login(c);upload(c)
    record=repo.list_all()[0]
    raw=json.loads(s3.objects[repo.ficha_key(record['id'])])
    assert raw['schema_version']==1 and 'extension' not in raw['archivo']
    assert from_record(raw,repo.prefix)==record
    streams=[]; old=s3.get_object
    def capture(**kw):
        obj=old(**kw);streams.append(obj['Body']);return obj
    s3.get_object=capture
    repo.get(record['id']);repo.download(record,1000)
    assert all(s.closed for s in streams)
    for op,kw in s3.calls:
        assert kw['IfNoneMatch']=='*' and kw['ServerSideEncryption']=='AES256'

@pytest.mark.parametrize('change',[
    lambda r:r.update(schema_version=True),lambda r:r.update(observaciones=None),
    lambda r:r.update(titulo=3),lambda r:r['archivo'].update(tamano_bytes=-1),
    lambda r:r['archivo'].update(tamano_bytes=True),lambda r:r['archivo'].update(sha256='bad'),
    lambda r:r['archivo'].update(clave_s3='other/original.txt'),
    lambda r:r['archivo'].update(clave_s3=r['archivo']['clave_s3']+'.exe'),
    lambda r:r['archivo'].update(tipo_mime='text/html'),
    lambda r:r['archivo'].update(nombre_original='../otro.txt'),
    lambda r:r.update(estado_declarado='aprobado'),lambda r:r['registro'].update(fecha='ayer'),
    lambda r:r['registro'].update(usuario=''),lambda r:r.update(relacion=[]),
])
def test_corrupt_record_fails_closed(env,change):
    _,c,s3,repo=env;login(c);upload(c)
    record=repo.list_all()[0];key=repo.ficha_key(record['id'])
    raw=json.loads(s3.objects[key]);change(raw);s3.objects[key]=json.dumps(raw).encode()
    assert c.get('/').status_code==503
    assert c.get('/programaciones/'+record['id']).status_code==503
    assert upload(c,content=b'New content').status_code==503


def test_oversized_record(env):
    _,c,s3,repo=env;login(c);upload(c)
    record=repo.list_all()[0];s3.objects[repo.ficha_key(record['id'])]=b' '* (128*1024+1)
    assert c.get('/').status_code==503


def test_same_hash_wrong_size_blocks_download(env):
    _,c,s3,repo=env;login(c);r=upload(c)
    record=repo.list_all()[0];key=repo.ficha_key(record['id']);raw=json.loads(s3.objects[key]);raw['archivo']['tamano_bytes']+=1;s3.objects[key]=json.dumps(raw).encode()
    assert c.get(r.location+'/descargar').status_code==503


def test_uncertain_ficha_keeps_original(env):
    _,c,s3,repo=env;login(c)
    s3.fail_ficha=True
    def fail_get(**kw): raise ClientError({'Error':{'Code':'AccessDenied'}},'GetObject')
    s3.get_object=fail_get
    r=upload(c)
    assert r.status_code==503 and b'original se conserva' in r.data
    assert len(s3.objects)==1


def test_failed_cleanup_is_reported(env):
    _,c,s3,_=env;login(c);s3.fail_ficha=True
    def fail_delete(**kw):raise ClientError({'Error':{'Code':'AccessDenied'}},'DeleteObject')
    s3.delete_object=fail_delete
    r=upload(c)
    assert r.status_code==503 and b'no se pudo retirar' in r.data
    assert len(s3.objects)==1


def test_original_precondition_never_deletes_existing(env):
    _,c,s3,repo=env;login(c);upload(c)
    record=repo.list_all()[0];before=deepcopy(s3.objects)
    with pytest.raises(StorageError):repo.save(record,b'different')
    assert s3.objects==before


def test_conflicting_ficha_never_deletes_original(env):
    _,c,s3,repo=env;login(c);upload(c)
    record=repo.list_all()[0]
    # Simulate existing different ficha at the generated UUID, no original yet.
    s3.objects.pop(original_key(repo,record))
    conflict=deepcopy(record);conflict['titulo']='Otro título conflictivo'
    s3.objects[repo.ficha_key(record['id'])]=json.dumps(to_record(conflict,repo.prefix)).encode()
    with pytest.raises(StorageError,match='Conflicto'):repo.save(record,b'Programacion de prueba para SBD')
    assert original_key(repo,record) in s3.objects
    assert repo.get(record['id'])['titulo']==conflict['titulo']


def test_lost_original_response_leaves_no_ficha(env):
    _,c,s3,_=env;login(c);old=s3.put_object
    def fail(**kw):
        old(**kw)
        raise ClientError({'Error':{'Code':'RequestTimeout'}},'PutObject')
    s3.put_object=fail
    assert upload(c).status_code==503
    assert len(s3.objects)==1 and all('/originales/' in k for k in s3.objects)


def test_bom_unchanged_and_filename_sanitized(env):
    _,c,_,repo=env;login(c);payload=b'\xef\xbb\xbfTexto de ejemplo\r\n'
    r=upload(c,content=payload,filename='carpeta/programacion.TXT')
    assert r.status_code==302
    record=repo.list_all()[0]
    assert record['archivo']['nombre_original']=='programacion.TXT'
    assert c.get(r.location+'/descargar').data==payload


def image_pdf():
    writer=PdfWriter();page=writer.add_blank_page(width=100,height=100)
    image=DecodedStreamObject();image.set_data(bytes([0,0,0]))
    image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(1),NameObject('/Height'):NumberObject(1),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im0'):writer._add_object(image)})})
    stream=DecodedStreamObject();stream.set_data(b'q 100 0 0 100 0 0 cm /Im0 Do Q')
    page[NameObject('/Contents')]=writer._add_object(stream)
    out=io.BytesIO();writer.write(out);return out.getvalue()


def test_real_image_pdf_accepted_without_ocr(env):
    _,c,_,_=env;login(c);payload=image_pdf()
    r=upload(c,content=payload,filename='imagen.pdf')
    assert r.status_code==302 and c.get(r.location+'/descargar').data==payload

@pytest.mark.parametrize('entry',['word/vbaProject.bin','word/embeddings/obj.bin'])
def test_docx_active_objects_rejected(env,entry):
    _,c,s3,_=env;login(c)
    out=io.BytesIO();Document().save(out)
    with zipfile.ZipFile(out,'a') as archive:archive.writestr(entry,b'not executable test content')
    assert upload(c,content=out.getvalue(),filename='objetos.docx').status_code==422
    assert not s3.objects


def test_docx_zip_entry_limit(env):
    _,c,s3,_=env;login(c);out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as archive:
        for i in range(2001):archive.writestr(str(i),'x')
    assert upload(c,content=out.getvalue(),filename='muchas.docx').status_code==422
    assert not s3.objects


def test_login_rate_limit_and_inactive_account(env):
    _,c,_,_=env;c.get('/login')
    with c.session_transaction() as session:token=session['csrf']
    for _ in range(10):
        assert c.post('/login',data=dict(csrf_token=token,username='docente',password='wrong')).status_code==401
    assert login(c).status_code==429


def test_invalid_unicode_csrf(env):
    _,c,_,_=env;c.get('/login')
    assert c.post('/login',data={'csrf_token':'no-válido'}).status_code==400


def test_arbitrary_relation_without_candidates(env):
    _,c,s3,_=env;login(c)
    assert upload(c,relacion_id=str(uuid4()),relacion_tipo='revision',relacion_motivo='Motivo válido').status_code==422
    assert not s3.objects


def test_windows_name_at_validator_boundary():
    file=DocumentFileValidator().validate(io.BytesIO(b'contenido'),r'C:\carpeta\programacion.TXT',100)
    assert file.filename=='programacion.TXT'


@pytest.mark.parametrize('entry,expected',[
    ('word/vbaProject.bin','declaración de macros'),
    ('word/embeddings/oleObject1.bin','no significa necesariamente que tenga macros'),
])
def test_docx_rejection_explains_actual_cause(entry,expected):
    out=io.BytesIO();Document().save(out)
    with zipfile.ZipFile(out,'a') as archive:archive.writestr(entry,b'synthetic object')
    with pytest.raises(ValidationError) as result:
        DocumentFileValidator().validate(io.BytesIO(out.getvalue()),'prueba.docx',1024*1024)
    assert expected in result.value.errors['archivo']


def test_empty_embeddings_directory_is_not_an_object():
    out=io.BytesIO();Document().save(out)
    with zipfile.ZipFile(out,'a') as archive:archive.writestr('word/embeddings/',b'')
    result=DocumentFileValidator().validate(io.BytesIO(out.getvalue()),'prueba.docx',1024*1024)
    assert result.content==out.getvalue()
